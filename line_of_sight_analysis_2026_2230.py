from pathlib import Path
import math

from qgis.core import (
    QgsApplication,
    QgsVectorLayer,
    QgsRasterLayer,
)

# =========================================================
# SETTINGS
# =========================================================

# Profile lines
LINE_PATH = Path(
    r"D:\Sali_IPRO\1.projects\3DVisualiyation\vector\SOL.gpkg"
)

# Current terrain: year 2026
DEM_2026_PATH = Path(
    r"D:\37_K+S_Nienburger_Mulde\260603_Landschaftsbild\dgm\32650-tiff-dgm2.tif"
)

# Future terrain after subsidence: year 2230
DEM_2230_PATH = Path(
    r"D:\Sali_IPRO\1.projects\3DVisualiyation\raster\DEM_2230.tif"
)

# New output directory
OUTPUT_DIR = Path(
    r"D:\Sali_IPRO\1.projects\3DVisualiyation\output\line_of_sight_analysis_2026_2230"
)

# Sampling interval along each profile line [m]
SAMPLE_INTERVAL = 2.0

# Observer eye height above the terrain [m]
CAMERA_HEIGHT = 1.8

# Line identification field
PNUMBER_FIELD = "PNumber"

# Display names
DEM_2026_NAME = "DEM 2026 - Current terrain"
DEM_2230_NAME = "DEM 2230 - Future terrain"


# =========================================================
# PLOT SETTINGS
# =========================================================

# Do not display beta values below this value.
# All values are nevertheless retained in Excel and used
# in the visibility calculation.
MIN_DISPLAY_BETA = -15.0

# Fixed range for Delta Beta plots
DELTA_Y_MIN = -1.0
DELTA_Y_MAX = 1.0


# =========================================================
# COLORS
# =========================================================

COLOR_2026 = "blue"
COLOR_2230 = "red"
COLOR_DELTA = "green"

# Numerical tolerance for the horizon test
ANGLE_TOLERANCE = 1e-9


# =========================================================
# INITIALIZE QGIS
# =========================================================

qgs = QgsApplication([], False)
qgs.initQgis()

try:

    # =====================================================
    # OUTPUT DIRECTORIES
    # =====================================================

    excel_dir = OUTPUT_DIR / "excel"
    viewing_angle_dir = OUTPUT_DIR / "01_viewing_angle"
    difference_dir = OUTPUT_DIR / "02_viewing_angle_difference"
    los_dir = OUTPUT_DIR / "03_line_of_sight"

    for folder in [
        excel_dir,
        viewing_angle_dir,
        difference_dir,
        los_dir,
    ]:

        folder.mkdir(
            parents=True,
            exist_ok=True
        )

    # =====================================================
    # CHECK INPUT FILES
    # =====================================================

    input_files = [
        (LINE_PATH, "Profile line layer"),
        (DEM_2026_PATH, "DEM 2026"),
        (DEM_2230_PATH, "DEM 2230"),
    ]

    for input_path, description in input_files:

        if not input_path.exists():

            raise FileNotFoundError(
                f"{description} was not found:\n"
                f"{input_path}"
            )

    # =====================================================
    # IMPORT MATPLOTLIB
    # =====================================================

    try:

        import matplotlib

        matplotlib.use("Agg")

        import matplotlib.pyplot as plt

    except ImportError:

        raise RuntimeError(
            "matplotlib is not available in the "
            "QGIS Python environment."
        )

    # =====================================================
    # IMPORT OPENPYXL
    # =====================================================

    try:

        from openpyxl import Workbook

        from openpyxl.styles import (
            Font,
            PatternFill,
            Alignment,
        )

        from openpyxl.utils import get_column_letter

    except ImportError:

        raise RuntimeError(
            "openpyxl is not available in the "
            "QGIS Python environment."
        )

    # =====================================================
    # LOAD INPUT DATA
    # =====================================================

    lines = QgsVectorLayer(
        str(LINE_PATH),
        "SOL",
        "ogr"
    )

    dem_2026 = QgsRasterLayer(
        str(DEM_2026_PATH),
        "DEM_2026"
    )

    dem_2230 = QgsRasterLayer(
        str(DEM_2230_PATH),
        "DEM_2230"
    )

    # =====================================================
    # VALIDATE INPUT LAYERS
    # =====================================================

    if not lines.isValid():

        raise RuntimeError(
            f"Could not load profile line layer:\n"
            f"{LINE_PATH}"
        )

    if not dem_2026.isValid():

        raise RuntimeError(
            f"Could not load DEM 2026:\n"
            f"{DEM_2026_PATH}"
        )

    if not dem_2230.isValid():

        raise RuntimeError(
            f"Could not load DEM 2230:\n"
            f"{DEM_2230_PATH}"
        )

    print("Input layers loaded successfully.")

    # =====================================================
    # CRS CHECK
    # =====================================================

    print(
        "Line CRS:",
        lines.crs().authid()
    )

    print(
        "DEM 2026 CRS:",
        dem_2026.crs().authid()
    )

    print(
        "DEM 2230 CRS:",
        dem_2230.crs().authid()
    )

    if not lines.crs().isValid():

        raise RuntimeError(
            "The CRS of the line layer is invalid."
        )

    if not dem_2026.crs().isValid():

        raise RuntimeError(
            "The CRS of DEM 2026 is invalid."
        )

    if not dem_2230.crs().isValid():

        raise RuntimeError(
            "The CRS of DEM 2230 is invalid."
        )

    if not (
        lines.crs() == dem_2026.crs()
        and lines.crs() == dem_2230.crs()
    ):

        raise RuntimeError(
            "CRS mismatch.\n"
            "The line layer, DEM 2026 and DEM 2230 "
            "must use the same CRS."
        )

    # Do not test mapUnits() as True/False.
    print(
        "Line layer map units:",
        lines.crs().mapUnits()
    )

    print(
        "CRS check completed successfully."
    )

    # =====================================================
    # FIELD CHECK
    # =====================================================

    available_fields = lines.fields().names()

    if PNUMBER_FIELD not in available_fields:

        raise RuntimeError(
            f"Field '{PNUMBER_FIELD}' was not found.\n"
            f"Available fields: {available_fields}"
        )

    # =====================================================
    # HELPER FUNCTIONS
    # =====================================================

    def sample_raster(raster, point):
        """
        Samples Band 1 of a raster.

        Returns None for invalid, NoData, NaN or
        infinite values.
        """

        value, ok = raster.dataProvider().sample(
            point,
            1
        )

        if not ok or value is None:
            return None

        try:
            value = float(value)

        except (TypeError, ValueError):
            return None

        if not math.isfinite(value):
            return None

        return value

    # -----------------------------------------------------

    def create_distances(length):
        """
        Creates sample distances:

        0, 2, 4, 6, ...

        The exact endpoint is added when necessary.
        """

        distances = []
        distance = 0.0

        while distance <= length:

            distances.append(
                distance
            )

            distance += SAMPLE_INTERVAL

        if (
            not distances
            or abs(distances[-1] - length) > 1e-8
        ):

            distances.append(
                length
            )

        return distances

    # -----------------------------------------------------

    def calculate_beta(
        terrain_elevation,
        eye_elevation,
        distance
    ):
        """
        Calculates the vertical viewing angle:

        beta = atan2(
            terrain elevation - eye elevation,
            horizontal distance
        )

        Positive beta:
            terrain is above the eye elevation.

        Negative beta:
            terrain is below the eye elevation.

        Beta at distance zero is undefined.
        """

        if terrain_elevation is None:
            return None

        if eye_elevation is None:
            return None

        if distance <= 0:
            return None

        return math.degrees(
            math.atan2(
                terrain_elevation - eye_elevation,
                distance
            )
        )

    # -----------------------------------------------------

    def calculate_visibility(
        rows,
        beta_key,
        visibility_key
    ):
        """
        Performs a one-dimensional terrain horizon test.

        A sample is visible when its beta angle is greater
        than the maximum beta angle of all preceding valid
        samples.

        All valid beta values are used, including values
        below MIN_DISPLAY_BETA.
        """

        maximum_previous_beta = -math.inf

        for row in rows:

            beta = row[beta_key]

            if beta is None:

                row[visibility_key] = None
                continue

            if maximum_previous_beta == -math.inf:

                visible = True

            else:

                visible = (
                    beta
                    > maximum_previous_beta
                    + ANGLE_TOLERANCE
                )

            row[visibility_key] = (
                "Visible"
                if visible
                else "Hidden"
            )

            if beta > maximum_previous_beta:

                maximum_previous_beta = beta

    # -----------------------------------------------------

    def format_excel(worksheet):

        header_fill = PatternFill(
            fill_type="solid",
            fgColor="D9EAF7"
        )

        for cell in worksheet[1]:

            cell.font = Font(
                bold=True
            )

            cell.fill = header_fill

            cell.alignment = Alignment(
                horizontal="center"
            )

        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions

        widths = {
            1: 12,
            2: 16,
            3: 30,
            4: 30,
            5: 27,
            6: 27,
            7: 33,
            8: 25,
            9: 25,
            10: 28,
        }

        for column, width in widths.items():

            worksheet.column_dimensions[
                get_column_letter(column)
            ].width = width

        # Numeric formatting for distance, elevation
        # and angle columns.
        for row in worksheet.iter_rows(
            min_row=2,
            min_col=2,
            max_col=7
        ):

            for cell in row:
                cell.number_format = "0.000"

    # -----------------------------------------------------

    def safe_filename(value):

        return (
            str(value)
            .strip()
            .replace(" ", "_")
            .replace("/", "_")
            .replace("\\", "_")
            .replace(":", "_")
        )

    # =====================================================
    # READ AND SORT FEATURES
    # =====================================================

    features = list(
        lines.getFeatures()
    )

    features.sort(
        key=lambda feature: (
            feature[PNUMBER_FIELD] is None,
            str(feature[PNUMBER_FIELD])
        )
    )

    print(
        f"Number of profile lines: {len(features)}"
    )

    # =====================================================
    # PROCESS EACH PROFILE
    # =====================================================

    processed_count = 0
    skipped_count = 0

    for feature in features:

        pnumber = feature[
            PNUMBER_FIELD
        ]

        geometry = feature.geometry()

        # -------------------------------------------------
        # PNUMBER CHECK
        # -------------------------------------------------

        if pnumber is None:

            print(
                "WARNING: A feature has no PNumber. "
                "Feature skipped."
            )

            skipped_count += 1
            continue

        # -------------------------------------------------
        # GEOMETRY CHECK
        # -------------------------------------------------

        if (
            geometry is None
            or geometry.isEmpty()
        ):

            print(
                f"WARNING: PNumber={pnumber}: "
                "empty geometry. Feature skipped."
            )

            skipped_count += 1
            continue

        length = geometry.length()

        if length <= 0:

            print(
                f"WARNING: PNumber={pnumber}: "
                "line length is zero. Feature skipped."
            )

            skipped_count += 1
            continue

        print()
        print("=" * 60)

        print(
            f"Processing PNumber={pnumber} | "
            f"Length={length:.2f} m"
        )

        # =================================================
        # CAMERA LOCATION
        #
        # The start point of the line is the observer
        # location.
        # =================================================

        camera_geometry = geometry.interpolate(
            0.0
        )

        if camera_geometry.isEmpty():

            print(
                "  ERROR: Camera location could not "
                "be determined."
            )

            skipped_count += 1
            continue

        camera_point = camera_geometry.asPoint()

        # =================================================
        # TERRAIN AT CAMERA
        # =================================================

        camera_ground_2026 = sample_raster(
            dem_2026,
            camera_point
        )

        camera_ground_2230 = sample_raster(
            dem_2230,
            camera_point
        )

        if camera_ground_2026 is None:

            print(
                "  ERROR: No valid DEM 2026 elevation "
                "at the camera location."
            )

            skipped_count += 1
            continue

        if camera_ground_2230 is None:

            print(
                "  ERROR: No valid DEM 2230 elevation "
                "at the camera location."
            )

            skipped_count += 1
            continue

        # =================================================
        # CAMERA / EYE ELEVATION
        #
        # Assumption:
        # In each period, the observer remains 1.8 m above
        # the terrain of that period. The observer therefore
        # moves vertically with terrain subsidence.
        # =================================================

        eye_elevation_2026 = (
            camera_ground_2026
            + CAMERA_HEIGHT
        )

        eye_elevation_2230 = (
            camera_ground_2230
            + CAMERA_HEIGHT
        )

        print(
            f"  DEM 2026 ground at camera: "
            f"{camera_ground_2026:.3f} m"
        )

        print(
            f"  DEM 2026 eye elevation: "
            f"{eye_elevation_2026:.3f} m"
        )

        print(
            f"  DEM 2230 ground at camera: "
            f"{camera_ground_2230:.3f} m"
        )

        print(
            f"  DEM 2230 eye elevation: "
            f"{eye_elevation_2230:.3f} m"
        )

        # =================================================
        # SAMPLE PROFILE
        # =================================================

        distances = create_distances(
            length
        )

        rows = []

        for distance in distances:

            point_geometry = geometry.interpolate(
                distance
            )

            if point_geometry.isEmpty():

                print(
                    f"  WARNING: Sample at "
                    f"{distance:.2f} m could not be created."
                )

                continue

            point = point_geometry.asPoint()

            # Current terrain
            elevation_2026 = sample_raster(
                dem_2026,
                point
            )

            # Future terrain
            elevation_2230 = sample_raster(
                dem_2230,
                point
            )

            # Current viewing angle
            beta_2026 = calculate_beta(
                elevation_2026,
                eye_elevation_2026,
                distance
            )

            # Future viewing angle
            beta_2230 = calculate_beta(
                elevation_2230,
                eye_elevation_2230,
                distance
            )

            # Future minus current
            delta_beta = None

            if (
                beta_2026 is not None
                and beta_2230 is not None
            ):

                delta_beta = (
                    beta_2230
                    - beta_2026
                )

            rows.append(
                {
                    "pnumber": pnumber,
                    "distance_m": distance,

                    "elevation_2026_m": elevation_2026,
                    "elevation_2230_m": elevation_2230,

                    "beta_2026_deg": beta_2026,
                    "beta_2230_deg": beta_2230,

                    "delta_beta_2230_minus_2026_deg": delta_beta,

                    "visibility_2026": None,
                    "visibility_2230": None,
                    "visibility_change_2026_to_2230": None,
                }
            )

        if not rows:

            print(
                f"  WARNING: No samples were created "
                f"for PNumber={pnumber}."
            )

            skipped_count += 1
            continue

        # =================================================
        # VISIBILITY CALCULATION
        # =================================================

        calculate_visibility(
            rows,
            "beta_2026_deg",
            "visibility_2026"
        )

        calculate_visibility(
            rows,
            "beta_2230_deg",
            "visibility_2230"
        )

        # =================================================
        # VISIBILITY CHANGE: CURRENT TO FUTURE
        # =================================================

        for row in rows:

            visibility_2026 = row[
                "visibility_2026"
            ]

            visibility_2230 = row[
                "visibility_2230"
            ]

            if (
                visibility_2026 is None
                or visibility_2230 is None
            ):

                row[
                    "visibility_change_2026_to_2230"
                ] = None

            else:

                row[
                    "visibility_change_2026_to_2230"
                ] = (
                    f"{visibility_2026} -> "
                    f"{visibility_2230}"
                )

        file_id = safe_filename(
            pnumber
        )

        # =================================================
        # CREATE EXCEL
        # =================================================

        workbook = Workbook()

        worksheet = workbook.active
        worksheet.title = "Line_of_Sight"

        worksheet.append(
            [
                "PNumber",
                "Distance_m",

                "Elevation_DEM_2026_Current_m",
                "Elevation_DEM_2230_Future_m",

                "Beta_DEM_2026_Current_deg",
                "Beta_DEM_2230_Future_deg",

                "Delta_Beta_2230_minus_2026_deg",

                "Visibility_DEM_2026_Current",
                "Visibility_DEM_2230_Future",

                "Visibility_Change_2026_to_2230",
            ]
        )

        for row in rows:

            worksheet.append(
                [
                    row["pnumber"],
                    row["distance_m"],

                    row["elevation_2026_m"],
                    row["elevation_2230_m"],

                    row["beta_2026_deg"],
                    row["beta_2230_deg"],

                    row[
                        "delta_beta_2230_minus_2026_deg"
                    ],

                    row["visibility_2026"],
                    row["visibility_2230"],

                    row[
                        "visibility_change_2026_to_2230"
                    ],
                ]
            )

        format_excel(
            worksheet
        )

        # =================================================
        # EXCEL VISIBILITY COLORS
        # =================================================

        fill_visible = PatternFill(
            fill_type="solid",
            fgColor="C6EFCE"
        )

        fill_hidden = PatternFill(
            fill_type="solid",
            fgColor="FFC7CE"
        )

        fill_lost = PatternFill(
            fill_type="solid",
            fgColor="F4CCCC"
        )

        fill_gained = PatternFill(
            fill_type="solid",
            fgColor="B7E1CD"
        )

        fill_unchanged = PatternFill(
            fill_type="solid",
            fgColor="E7E6E6"
        )

        for row_number in range(
            2,
            worksheet.max_row + 1
        ):

            # Visibility columns
            for column_number in [8, 9]:

                cell = worksheet.cell(
                    row=row_number,
                    column=column_number
                )

                if cell.value == "Visible":

                    cell.fill = fill_visible

                elif cell.value == "Hidden":

                    cell.fill = fill_hidden

            # Visibility-change column
            change_cell = worksheet.cell(
                row=row_number,
                column=10
            )

            if change_cell.value == "Visible -> Hidden":

                change_cell.fill = fill_lost

            elif change_cell.value == "Hidden -> Visible":

                change_cell.fill = fill_gained

            elif change_cell.value in [
                "Visible -> Visible",
                "Hidden -> Hidden",
            ]:

                change_cell.fill = fill_unchanged

        # =================================================
        # PARAMETERS SHEET
        # =================================================

        parameters_sheet = workbook.create_sheet(
            "Parameters"
        )

        parameters = [
            ["Parameter", "Value"],

            ["PNumber", pnumber],

            ["CRS", lines.crs().authid()],

            ["Line_Length_m", length],

            ["Sample_Interval_m", SAMPLE_INTERVAL],

            [
                "Camera_Height_Above_Terrain_m",
                CAMERA_HEIGHT
            ],

            [
                "DEM_2026_Path",
                str(DEM_2026_PATH)
            ],

            [
                "DEM_2230_Path",
                str(DEM_2230_PATH)
            ],

            [
                "DEM_2026_Ground_At_Camera_m",
                camera_ground_2026
            ],

            [
                "DEM_2026_Eye_Elevation_m",
                eye_elevation_2026
            ],

            [
                "DEM_2230_Ground_At_Camera_m",
                camera_ground_2230
            ],

            [
                "DEM_2230_Eye_Elevation_m",
                eye_elevation_2230
            ],

            [
                "Camera_Assumption",
                (
                    "Observer is 1.8 m above the terrain "
                    "in each period and moves vertically "
                    "with terrain subsidence."
                )
            ],

            [
                "Beta_Definition",
                (
                    "atan2("
                    "terrain elevation - eye elevation, "
                    "horizontal distance)"
                )
            ],

            [
                "Delta_Beta_Definition",
                "Beta_DEM_2230 - Beta_DEM_2026"
            ],

            [
                "Visibility_Change_Direction",
                "DEM 2026 -> DEM 2230"
            ],

            [
                "Viewing_Angle_Display_Min_deg",
                MIN_DISPLAY_BETA
            ],

            [
                "Delta_Plot_Y_Min_deg",
                DELTA_Y_MIN
            ],

            [
                "Delta_Plot_Y_Max_deg",
                DELTA_Y_MAX
            ],

            [
                "Visibility_Method",
                "1D terrain-profile horizon test"
            ],

            [
                "Beta_Below_Display_Min_Used_For_Visibility",
                "Yes"
            ],

            [
                "Earth_Curvature_And_Refraction",
                "Not included"
            ],
        ]

        for item in parameters:

            parameters_sheet.append(
                item
            )

        parameters_sheet["A1"].font = Font(
            bold=True
        )

        parameters_sheet["B1"].font = Font(
            bold=True
        )

        parameters_sheet.column_dimensions["A"].width = 50
        parameters_sheet.column_dimensions["B"].width = 100
        parameters_sheet.freeze_panes = "A2"

        excel_path = (
            excel_dir
            / f"line_of_sight_{file_id}.xlsx"
        )

        workbook.save(
            excel_path
        )

        # =================================================
        # PLOT 1: VIEWING ANGLES
        # =================================================

        display_2026 = [
            row
            for row in rows
            if (
                row["beta_2026_deg"] is not None
                and row["beta_2026_deg"] >= MIN_DISPLAY_BETA
            )
        ]

        display_2230 = [
            row
            for row in rows
            if (
                row["beta_2230_deg"] is not None
                and row["beta_2230_deg"] >= MIN_DISPLAY_BETA
            )
        ]

        x_2026 = [
            row["distance_m"]
            for row in display_2026
        ]

        y_2026 = [
            row["beta_2026_deg"]
            for row in display_2026
        ]

        x_2230 = [
            row["distance_m"]
            for row in display_2230
        ]

        y_2230 = [
            row["beta_2230_deg"]
            for row in display_2230
        ]

        fig, axis = plt.subplots(
            figsize=(13, 6)
        )

        if x_2026:

            axis.plot(
                x_2026,
                y_2026,
                color=COLOR_2026,
                linewidth=1.5,
                label=DEM_2026_NAME
            )

        if x_2230:

            axis.plot(
                x_2230,
                y_2230,
                color=COLOR_2230,
                linewidth=1.5,
                label=DEM_2230_NAME
            )

        axis.axhline(
            0,
            color="black",
            linestyle="--",
            linewidth=1,
            alpha=0.5
        )

        axis.set_xlim(
            0,
            length
        )

        axis.set_ylim(
            bottom=MIN_DISPLAY_BETA
        )

        axis.set_title(
            f"Viewing-Angle Profile - PNumber {pnumber}"
        )

        axis.set_xlabel(
            "Distance from camera [m]"
        )

        axis.set_ylabel(
            "Vertical viewing angle beta [deg]"
        )

        axis.grid(
            True,
            alpha=0.25
        )

        if x_2026 or x_2230:

            axis.legend(
                loc="best"
            )

        fig.tight_layout()

        viewing_path = (
            viewing_angle_dir
            / f"viewing_angle_{file_id}.png"
        )

        fig.savefig(
            viewing_path,
            dpi=200,
            bbox_inches="tight"
        )

        plt.close(
            fig
        )

        # =================================================
        # PLOT 2: DELTA BETA
        #
        # Delta Beta = Beta 2230 - Beta 2026
        # =================================================

        delta_rows = [
            row
            for row in rows
            if row[
                "delta_beta_2230_minus_2026_deg"
            ] is not None
        ]

        delta_x = [
            row["distance_m"]
            for row in delta_rows
        ]

        delta_y = [
            row[
                "delta_beta_2230_minus_2026_deg"
            ]
            for row in delta_rows
        ]

        fig, axis = plt.subplots(
            figsize=(13, 5.5)
        )

        if delta_x:

            axis.plot(
                delta_x,
                delta_y,
                color=COLOR_DELTA,
                linewidth=1.5,
                label="Delta beta = Beta 2230 - Beta 2026"
            )

        axis.axhline(
            0,
            color="black",
            linestyle="--",
            linewidth=1,
            alpha=0.6
        )

        axis.set_xlim(
            0,
            length
        )

        axis.set_ylim(
            DELTA_Y_MIN,
            DELTA_Y_MAX
        )

        axis.set_title(
            f"Viewing-Angle Change - PNumber {pnumber}"
        )

        axis.set_xlabel(
            "Distance from camera [m]"
        )

        axis.set_ylabel(
            "Delta beta: 2230 - 2026 [deg]"
        )

        axis.grid(
            True,
            alpha=0.25
        )

        if delta_x:

            axis.legend(
                loc="best"
            )

        fig.tight_layout()

        difference_path = (
            difference_dir
            / f"viewing_angle_change_{file_id}.png"
        )

        fig.savefig(
            difference_path,
            dpi=200,
            bbox_inches="tight"
        )

        plt.close(
            fig
        )

        # =================================================
        # PLOT 3: LINE OF SIGHT
        # =================================================

        los_2026 = display_2026
        los_2230 = display_2230

        visible_2026 = [
            row
            for row in los_2026
            if row["visibility_2026"] == "Visible"
        ]

        hidden_2026 = [
            row
            for row in los_2026
            if row["visibility_2026"] == "Hidden"
        ]

        visible_2230 = [
            row
            for row in los_2230
            if row["visibility_2230"] == "Visible"
        ]

        hidden_2230 = [
            row
            for row in los_2230
            if row["visibility_2230"] == "Hidden"
        ]

        fig, axis = plt.subplots(
            figsize=(15, 7)
        )

        # Background angle profiles
        if los_2026:

            axis.plot(
                [
                    row["distance_m"]
                    for row in los_2026
                ],
                [
                    row["beta_2026_deg"]
                    for row in los_2026
                ],
                color=COLOR_2026,
                linewidth=0.8,
                alpha=0.25
            )

        if los_2230:

            axis.plot(
                [
                    row["distance_m"]
                    for row in los_2230
                ],
                [
                    row["beta_2230_deg"]
                    for row in los_2230
                ],
                color=COLOR_2230,
                linewidth=0.8,
                alpha=0.25
            )

        # DEM 2026 visible
        axis.scatter(
            [
                row["distance_m"]
                for row in visible_2026
            ],
            [
                row["beta_2026_deg"]
                for row in visible_2026
            ],
            color=COLOR_2026,
            marker="o",
            s=24,
            alpha=0.85,
            edgecolors="none",
            label="DEM 2026 Visible",
            zorder=4
        )

        # DEM 2026 hidden
        axis.scatter(
            [
                row["distance_m"]
                for row in hidden_2026
            ],
            [
                row["beta_2026_deg"]
                for row in hidden_2026
            ],
            color=COLOR_2026,
            marker="x",
            s=20,
            linewidths=0.8,
            alpha=0.38,
            label="DEM 2026 Hidden",
            zorder=3
        )

        # DEM 2230 visible
        axis.scatter(
            [
                row["distance_m"]
                for row in visible_2230
            ],
            [
                row["beta_2230_deg"]
                for row in visible_2230
            ],
            color=COLOR_2230,
            marker="o",
            s=24,
            alpha=0.85,
            edgecolors="none",
            label="DEM 2230 Visible",
            zorder=4
        )

        # DEM 2230 hidden
        axis.scatter(
            [
                row["distance_m"]
                for row in hidden_2230
            ],
            [
                row["beta_2230_deg"]
                for row in hidden_2230
            ],
            color=COLOR_2230,
            marker="x",
            s=20,
            linewidths=0.8,
            alpha=0.38,
            label="DEM 2230 Hidden",
            zorder=3
        )

        axis.axhline(
            0,
            color="black",
            linestyle="--",
            linewidth=1,
            alpha=0.5
        )

        axis.set_xlim(
            0,
            length
        )

        axis.set_ylim(
            bottom=MIN_DISPLAY_BETA
        )

        axis.set_title(
            f"Line-of-Sight Comparison - PNumber {pnumber}",
            fontsize=14
        )

        axis.set_xlabel(
            "Distance from camera [m]"
        )

        axis.set_ylabel(
            "Vertical viewing angle beta [deg]"
        )

        axis.grid(
            True,
            alpha=0.20
        )

        axis.legend(
            loc="best",
            fontsize=9,
            framealpha=0.90
        )

        fig.tight_layout()

        los_path = (
            los_dir
            / f"line_of_sight_{file_id}.png"
        )

        fig.savefig(
            los_path,
            dpi=220,
            bbox_inches="tight"
        )

        plt.close(
            fig
        )

        # =================================================
        # SUMMARY
        # =================================================

        analysis_rows = [
            row
            for row in rows
            if row["distance_m"] > 0
        ]

        visible_2026_count = sum(
            row["visibility_2026"] == "Visible"
            for row in analysis_rows
        )

        hidden_2026_count = sum(
            row["visibility_2026"] == "Hidden"
            for row in analysis_rows
        )

        visible_2230_count = sum(
            row["visibility_2230"] == "Visible"
            for row in analysis_rows
        )

        hidden_2230_count = sum(
            row["visibility_2230"] == "Hidden"
            for row in analysis_rows
        )

        visibility_lost_count = sum(
            row[
                "visibility_change_2026_to_2230"
            ] == "Visible -> Hidden"
            for row in analysis_rows
        )

        visibility_gained_count = sum(
            row[
                "visibility_change_2026_to_2230"
            ] == "Hidden -> Visible"
            for row in analysis_rows
        )

        print(
            f"  Samples: {len(rows)}"
        )

        print(
            f"  DEM 2026: "
            f"{visible_2026_count} visible / "
            f"{hidden_2026_count} hidden"
        )

        print(
            f"  DEM 2230: "
            f"{visible_2230_count} visible / "
            f"{hidden_2230_count} hidden"
        )

        print(
            f"  Visibility lost "
            f"(Visible -> Hidden): "
            f"{visibility_lost_count}"
        )

        print(
            f"  Visibility gained "
            f"(Hidden -> Visible): "
            f"{visibility_gained_count}"
        )

        print(
            f"  Excel: {excel_path.name}"
        )

        print(
            f"  Viewing Angle: {viewing_path.name}"
        )

        print(
            f"  Difference: {difference_path.name}"
        )

        print(
            f"  Line of Sight: {los_path.name}"
        )

        processed_count += 1

    # =====================================================
    # FINISHED
    # =====================================================

    print()
    print("=" * 60)
    print("LINE-OF-SIGHT PROCESSING COMPLETED")
    print("=" * 60)

    print(
        f"Processed lines: {processed_count}"
    )

    print(
        f"Skipped lines: {skipped_count}"
    )

    print(
        f"Output:\n{OUTPUT_DIR}"
    )

    print()

    print(
        "Delta Beta = Beta 2230 - Beta 2026"
    )

    print(
        "Visibility change = Visibility 2026 -> Visibility 2230"
    )

    print(
        f"Viewing Angle / LOS display: "
        f"beta >= {MIN_DISPLAY_BETA} deg"
    )

    print(
        f"Delta Beta display: "
        f"{DELTA_Y_MIN} to {DELTA_Y_MAX} deg"
    )

    print(
        "Values outside plot limits remain unchanged "
        "in Excel and are used in the visibility analysis."
    )

finally:

    qgs.exitQgis()