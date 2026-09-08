from pathlib import Path
import math

from qgis.core import (
    QgsApplication,
    QgsVectorLayer,
    QgsRasterLayer,
)

# =========================================================
# SETTINGS
# =========================================================

# Profile lines
LINE_PATH = Path(
    r"D:\Sali_IPRO\1.projects\3DVisualiyation\vector\SOL.gpkg"
)

# Current terrain: year 2026
DEM_2026_PATH = Path(
    r"D:\37_K+S_Nienburger_Mulde\260603_Landschaftsbild\dgm\32650-tiff-dgm2.tif"
)

# Future terrain: year 2230
DEM_2230_PATH = Path(
    r"D:\Sali_IPRO\1.projects\3DVisualiyation\raster\DEM_2230.tif"
)

# New output directory
OUTPUT_DIR = Path(
    r"D:\Sali_IPRO\1.projects\3DVisualiyation\output\viewing_angle_profiles_2026_2230"
)

# Sampling interval along each profile line [m]
SAMPLE_INTERVAL = 2.0

# Camera/eye height above the terrain of each period [m]
CAMERA_HEIGHT = 1.8

# Line identification field
PNUMBER_FIELD = "PNumber"

# Names used in plots and output files
DEM_2026_NAME = "DEM 2026 - Current terrain"
DEM_2230_NAME = "DEM 2230 - Future terrain"

# Plot colors
COLOR_2026 = "blue"
COLOR_2230 = "red"
COLOR_DELTA = "green"


# =========================================================
# INITIALIZE QGIS
# =========================================================

qgs = QgsApplication([], False)
qgs.initQgis()

try:

    # =====================================================
    # CREATE OUTPUT DIRECTORIES
    # =====================================================

    excel_dir = OUTPUT_DIR / "excel"
    plot_comparison_dir = OUTPUT_DIR / "01_beta_comparison"
    plot_difference_dir = OUTPUT_DIR / "02_beta_with_difference"

    excel_dir.mkdir(
        parents=True,
        exist_ok=True
    )

    plot_comparison_dir.mkdir(
        parents=True,
        exist_ok=True
    )

    plot_difference_dir.mkdir(
        parents=True,
        exist_ok=True
    )

    # =====================================================
    # CHECK INPUT FILES
    # =====================================================

    input_files = [
        (LINE_PATH, "Profile line layer"),
        (DEM_2026_PATH, "DEM 2026"),
        (DEM_2230_PATH, "DEM 2230"),
    ]

    for input_path, input_description in input_files:

        if not input_path.exists():

            raise FileNotFoundError(
                f"{input_description} was not found:\n"
                f"{input_path}"
            )

    # =====================================================
    # IMPORT MATPLOTLIB
    # =====================================================

    try:

        import matplotlib

        # Required for execution without a GUI
        matplotlib.use("Agg")

        import matplotlib.pyplot as plt

    except ImportError:

        raise RuntimeError(
            "matplotlib is not available in the "
            "QGIS Python environment."
        )

    # =====================================================
    # IMPORT OPENPYXL
    # =====================================================

    try:

        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter

    except ImportError:

        raise RuntimeError(
            "openpyxl is not available in the "
            "QGIS Python environment.\n"
            "It is required to create XLSX files."
        )

    # =====================================================
    # LOAD INPUT DATA
    # =====================================================

    lines = QgsVectorLayer(
        str(LINE_PATH),
        "SOL",
        "ogr"
    )

    dem_2026 = QgsRasterLayer(
        str(DEM_2026_PATH),
        "DEM_2026"
    )

    dem_2230 = QgsRasterLayer(
        str(DEM_2230_PATH),
        "DEM_2230"
    )

    # =====================================================
    # VALIDATE INPUT LAYERS
    # =====================================================

    if not lines.isValid():

        raise RuntimeError(
            f"Could not load profile line layer:\n"
            f"{LINE_PATH}"
        )

    if not dem_2026.isValid():

        raise RuntimeError(
            f"Could not load DEM 2026:\n"
            f"{DEM_2026_PATH}"
        )

    if not dem_2230.isValid():

        raise RuntimeError(
            f"Could not load DEM 2230:\n"
            f"{DEM_2230_PATH}"
        )

    print("Input layers loaded successfully.")

    # =====================================================
    # CRS CHECK
    # =====================================================

    print(
        "Line CRS:",
        lines.crs().authid()
    )

    print(
        "DEM 2026 CRS:",
        dem_2026.crs().authid()
    )

    print(
        "DEM 2230 CRS:",
        dem_2230.crs().authid()
    )

    if not lines.crs().isValid():

        raise RuntimeError(
            "The CRS of the profile line layer is invalid."
        )

    if not dem_2026.crs().isValid():

        raise RuntimeError(
            "The CRS of DEM 2026 is invalid."
        )

    if not dem_2230.crs().isValid():

        raise RuntimeError(
            "The CRS of DEM 2230 is invalid."
        )

    if not (
        lines.crs() == dem_2026.crs()
        and lines.crs() == dem_2230.crs()
    ):

        raise RuntimeError(
            "CRS mismatch.\n"
            "The profile line layer, DEM 2026 and DEM 2230 "
            "must use the same CRS."
        )

    # Only print the unit. Do not test it as True/False.
    print(
        "Line layer map units:",
        lines.crs().mapUnits()
    )

    print(
        "CRS check completed successfully."
    )

    # EPSG:25832 is a projected CRS in meters.
    # Therefore SAMPLE_INTERVAL = 2.0 means 2 meters.

    # =====================================================
    # FIELD CHECK
    # =====================================================

    available_fields = (
        lines.fields().names()
    )

    if PNUMBER_FIELD not in available_fields:

        raise RuntimeError(
            f"Field '{PNUMBER_FIELD}' was not found.\n"
            f"Available fields: {available_fields}"
        )

    # =====================================================
    # HELPER FUNCTION: SAMPLE RASTER
    # =====================================================

    def sample_raster(raster, point):
        """
        Samples Band 1 of a raster at the specified point.

        Returns None for invalid, NoData, nonnumeric,
        NaN or infinite values.
        """

        value, ok = (
            raster
            .dataProvider()
            .sample(
                point,
                1
            )
        )

        if not ok or value is None:
            return None

        try:
            value = float(value)

        except (TypeError, ValueError):
            return None

        if not math.isfinite(value):
            return None

        return value

    # =====================================================
    # HELPER FUNCTION: CREATE SAMPLE DISTANCES
    # =====================================================

    def create_distances(length):
        """
        Creates sample distances:

        0, 2, 4, 6, ...

        The exact endpoint is added when necessary.
        """

        distances = []
        distance = 0.0

        while distance <= length:

            distances.append(
                distance
            )

            distance += SAMPLE_INTERVAL

        if (
            not distances
            or abs(distances[-1] - length) > 1e-8
        ):

            distances.append(
                length
            )

        return distances

    # =====================================================
    # HELPER FUNCTION: CALCULATE VIEWING ANGLE
    # =====================================================

    def calculate_beta(
        terrain_elevation,
        eye_elevation,
        distance
    ):
        """
        Calculates the vertical viewing angle:

        beta = atan2(
            terrain elevation - eye elevation,
            horizontal distance
        )

        The result is returned in degrees.

        Positive beta:
            target terrain is above the eye elevation.

        Negative beta:
            target terrain is below the eye elevation.
        """

        if terrain_elevation is None:
            return None

        if eye_elevation is None:
            return None

        # The viewing angle at the camera position itself
        # is undefined and is therefore stored as None.
        if distance <= 0:
            return None

        return math.degrees(
            math.atan2(
                terrain_elevation - eye_elevation,
                distance
            )
        )

    # =====================================================
    # HELPER FUNCTION: FORMAT EXCEL
    # =====================================================

    def format_excel_sheet(worksheet):

        header_fill = PatternFill(
            fill_type="solid",
            fgColor="D9EAF7"
        )

        for cell in worksheet[1]:

            cell.font = Font(
                bold=True
            )

            cell.fill = header_fill

            cell.alignment = Alignment(
                horizontal="center"
            )

        widths = {
            1: 12,
            2: 16,
            3: 28,
            4: 28,
            5: 24,
            6: 24,
            7: 28,
        }

        for column_number, width in widths.items():

            worksheet.column_dimensions[
                get_column_letter(column_number)
            ].width = width

        worksheet.freeze_panes = "A2"

        for row in worksheet.iter_rows(
            min_row=2,
            min_col=2,
            max_col=7
        ):

            for cell in row:

                cell.number_format = "0.000"

    # =====================================================
    # HELPER FUNCTION: SAFE FILE NAME
    # =====================================================

    def safe_filename(value):
        """
        Converts a PNumber value into a safe filename.
        """

        return (
            str(value)
            .strip()
            .replace(" ", "_")
            .replace("/", "_")
            .replace("\\", "_")
            .replace(":", "_")
        )

    # =====================================================
    # READ AND SORT FEATURES
    # =====================================================

    features = list(
        lines.getFeatures()
    )

    features.sort(
        key=lambda feature: (
            feature[PNUMBER_FIELD] is None,
            str(feature[PNUMBER_FIELD])
        )
    )

    print(
        f"Number of profile lines: {len(features)}"
    )

    print()

    # =====================================================
    # PROCESS EACH PROFILE LINE
    # =====================================================

    processed_count = 0
    skipped_count = 0

    for feature in features:

        pnumber = feature[
            PNUMBER_FIELD
        ]

        geometry = feature.geometry()

        # -------------------------------------------------
        # CHECK PNUMBER
        # -------------------------------------------------

        if pnumber is None:

            print(
                "WARNING: A feature has no PNumber. "
                "Feature skipped."
            )

            skipped_count += 1
            continue

        # -------------------------------------------------
        # CHECK GEOMETRY
        # -------------------------------------------------

        if (
            geometry is None
            or geometry.isEmpty()
        ):

            print(
                f"WARNING: PNumber={pnumber}: "
                "empty geometry. Feature skipped."
            )

            skipped_count += 1
            continue

        length = geometry.length()

        if length <= 0:

            print(
                f"WARNING: PNumber={pnumber}: "
                "line length is zero. Feature skipped."
            )

            skipped_count += 1
            continue

        print("=" * 60)

        print(
            f"Processing PNumber={pnumber} | "
            f"Length={length:.2f} m"
        )

        # =================================================
        # CAMERA POSITION
        #
        # The start of each profile line is treated as
        # the camera location.
        # =================================================

        camera_geometry = geometry.interpolate(
            0.0
        )

        if camera_geometry.isEmpty():

            print(
                "  ERROR: Camera location could not be determined."
            )

            skipped_count += 1
            continue

        camera_point = (
            camera_geometry.asPoint()
        )

        # Terrain elevation at camera location in 2026
        camera_ground_2026 = sample_raster(
            dem_2026,
            camera_point
        )

        # Terrain elevation at camera location in 2230
        camera_ground_2230 = sample_raster(
            dem_2230,
            camera_point
        )

        if camera_ground_2026 is None:

            print(
                "  ERROR: DEM 2026 has no valid elevation "
                "at the camera location."
            )

            skipped_count += 1
            continue

        if camera_ground_2230 is None:

            print(
                "  ERROR: DEM 2230 has no valid elevation "
                "at the camera location."
            )

            skipped_count += 1
            continue

        # =================================================
        # EYE ELEVATIONS
        #
        # Assumption:
        # The observer is 1.8 m above the terrain in each
        # period. Therefore the observer moves vertically
        # with the subsiding terrain.
        # =================================================

        eye_elevation_2026 = (
            camera_ground_2026
            + CAMERA_HEIGHT
        )

        eye_elevation_2230 = (
            camera_ground_2230
            + CAMERA_HEIGHT
        )

        print(
            f"  DEM 2026 terrain at camera: "
            f"{camera_ground_2026:.3f} m"
        )

        print(
            f"  DEM 2026 eye elevation: "
            f"{eye_elevation_2026:.3f} m"
        )

        print(
            f"  DEM 2230 terrain at camera: "
            f"{camera_ground_2230:.3f} m"
        )

        print(
            f"  DEM 2230 eye elevation: "
            f"{eye_elevation_2230:.3f} m"
        )

        # =================================================
        # SAMPLE PROFILE AND CALCULATE BETA
        # =================================================

        distances = create_distances(
            length
        )

        rows = []

        for distance in distances:

            point_geometry = geometry.interpolate(
                distance
            )

            if point_geometry.isEmpty():

                print(
                    f"  WARNING: Could not create sample "
                    f"at distance {distance:.2f} m."
                )

                continue

            point = point_geometry.asPoint()

            # Current terrain elevation
            elevation_2026 = sample_raster(
                dem_2026,
                point
            )

            # Future terrain elevation
            elevation_2230 = sample_raster(
                dem_2230,
                point
            )

            # Current viewing angle
            beta_2026 = calculate_beta(
                elevation_2026,
                eye_elevation_2026,
                distance
            )

            # Future viewing angle
            beta_2230 = calculate_beta(
                elevation_2230,
                eye_elevation_2230,
                distance
            )

            # ---------------------------------------------
            # VIEWING-ANGLE CHANGE
            #
            # Future minus current:
            #
            # Delta Beta = Beta 2230 - Beta 2026
            #
            # Positive:
            # The viewing angle is higher in 2230.
            #
            # Negative:
            # The viewing angle is lower in 2230.
            # ---------------------------------------------

            delta_beta = None

            if (
                beta_2026 is not None
                and beta_2230 is not None
            ):

                delta_beta = (
                    beta_2230
                    - beta_2026
                )

            rows.append(
                {
                    "pnumber": pnumber,
                    "distance_m": distance,
                    "elevation_2026_m": elevation_2026,
                    "elevation_2230_m": elevation_2230,
                    "beta_2026_deg": beta_2026,
                    "beta_2230_deg": beta_2230,
                    "delta_beta_2230_minus_2026_deg": delta_beta,
                }
            )

        if not rows:

            print(
                f"  WARNING: No samples were created "
                f"for PNumber={pnumber}."
            )

            skipped_count += 1
            continue

        filename_id = safe_filename(
            pnumber
        )

        # =================================================
        # CREATE EXCEL FILE
        # =================================================

        workbook = Workbook()

        worksheet = workbook.active
        worksheet.title = "Viewing_Angle"

        headers = [
            "PNumber",
            "Distance_m",
            "Elevation_DEM_2026_Current_m",
            "Elevation_DEM_2230_Future_m",
            "Beta_DEM_2026_Current_deg",
            "Beta_DEM_2230_Future_deg",
            "Delta_Beta_2230_minus_2026_deg",
        ]

        worksheet.append(
            headers
        )

        for row in rows:

            worksheet.append(
                [
                    row["pnumber"],
                    row["distance_m"],
                    row["elevation_2026_m"],
                    row["elevation_2230_m"],
                    row["beta_2026_deg"],
                    row["beta_2230_deg"],
                    row["delta_beta_2230_minus_2026_deg"],
                ]
            )

        format_excel_sheet(
            worksheet
        )

        # =================================================
        # PARAMETERS / METADATA SHEET
        # =================================================

        metadata_sheet = workbook.create_sheet(
            "Parameters"
        )

        metadata = [
            ["Parameter", "Value"],
            ["PNumber", pnumber],
            ["CRS", lines.crs().authid()],
            ["Line_Length_m", length],
            ["Sample_Interval_m", SAMPLE_INTERVAL],
            ["Camera_Height_Above_Terrain_m", CAMERA_HEIGHT],
            ["DEM_2026_Path", str(DEM_2026_PATH)],
            ["DEM_2230_Path", str(DEM_2230_PATH)],
            [
                "DEM_2026_Terrain_At_Camera_m",
                camera_ground_2026
            ],
            [
                "DEM_2026_Eye_Elevation_m",
                eye_elevation_2026
            ],
            [
                "DEM_2230_Terrain_At_Camera_m",
                camera_ground_2230
            ],
            [
                "DEM_2230_Eye_Elevation_m",
                eye_elevation_2230
            ],
            [
                "Beta_Definition",
                "atan2(zTerrain - zEye, distance)"
            ],
            [
                "Delta_Beta_Definition",
                "Beta_DEM_2230 - Beta_DEM_2026"
            ],
            [
                "Camera_Assumption",
                (
                    "Observer is 1.8 m above the terrain "
                    "in each period and moves with terrain subsidence."
                )
            ],
        ]

        for item in metadata:

            metadata_sheet.append(
                item
            )

        metadata_sheet["A1"].font = Font(
            bold=True
        )

        metadata_sheet["B1"].font = Font(
            bold=True
        )

        metadata_sheet.column_dimensions["A"].width = 42
        metadata_sheet.column_dimensions["B"].width = 95

        metadata_sheet.freeze_panes = "A2"

        excel_path = (
            excel_dir
            / f"viewing_angle_{filename_id}.xlsx"
        )

        workbook.save(
            excel_path
        )

        # =================================================
        # PREPARE PLOT DATA
        # =================================================

        x_2026 = [
            row["distance_m"]
            for row in rows
            if row["beta_2026_deg"] is not None
        ]

        beta_2026_values = [
            row["beta_2026_deg"]
            for row in rows
            if row["beta_2026_deg"] is not None
        ]

        x_2230 = [
            row["distance_m"]
            for row in rows
            if row["beta_2230_deg"] is not None
        ]

        beta_2230_values = [
            row["beta_2230_deg"]
            for row in rows
            if row["beta_2230_deg"] is not None
        ]

        x_delta = [
            row["distance_m"]
            for row in rows
            if row[
                "delta_beta_2230_minus_2026_deg"
            ] is not None
        ]

        delta_beta_values = [
            row[
                "delta_beta_2230_minus_2026_deg"
            ]
            for row in rows
            if row[
                "delta_beta_2230_minus_2026_deg"
            ] is not None
        ]

        # =================================================
        # PLOT 1: COMPARISON OF VIEWING ANGLES
        # =================================================

        fig, axis = plt.subplots(
            figsize=(12, 6)
        )

        if x_2026:

            axis.plot(
                x_2026,
                beta_2026_values,
                color=COLOR_2026,
                linewidth=1.5,
                label=DEM_2026_NAME
            )

        if x_2230:

            axis.plot(
                x_2230,
                beta_2230_values,
                color=COLOR_2230,
                linewidth=1.5,
                label=DEM_2230_NAME
            )

        # Horizontal viewing direction
        axis.axhline(
            y=0,
            color="black",
            linewidth=1.0,
            linestyle="--",
            alpha=0.60
        )

        axis.set_xlim(
            0,
            length
        )

        axis.set_title(
            f"Viewing-Angle Profile - PNumber {pnumber}",
            fontsize=14
        )

        axis.set_xlabel(
            "Distance from camera [m]",
            fontsize=11
        )

        axis.set_ylabel(
            "Vertical viewing angle beta [deg]",
            fontsize=11
        )

        axis.grid(
            True,
            alpha=0.30
        )

        if x_2026 or x_2230:

            axis.legend(
                loc="best"
            )

        fig.tight_layout()

        comparison_plot_path = (
            plot_comparison_dir
            / f"viewing_angle_{filename_id}.png"
        )

        fig.savefig(
            comparison_plot_path,
            dpi=200,
            bbox_inches="tight"
        )

        plt.close(
            fig
        )

        # =================================================
        # PLOT 2: BETAS AND DELTA BETA
        #
        # Left Y-axis:
        # Beta 2026 and Beta 2230
        #
        # Right Y-axis:
        # Delta Beta = Beta 2230 - Beta 2026
        # =================================================

        fig, axis_beta = plt.subplots(
            figsize=(12, 6)
        )

        line_handles = []

        if x_2026:

            line_2026, = axis_beta.plot(
                x_2026,
                beta_2026_values,
                color=COLOR_2026,
                linewidth=1.5,
                label=DEM_2026_NAME
            )

            line_handles.append(
                line_2026
            )

        if x_2230:

            line_2230, = axis_beta.plot(
                x_2230,
                beta_2230_values,
                color=COLOR_2230,
                linewidth=1.5,
                label=DEM_2230_NAME
            )

            line_handles.append(
                line_2230
            )

        axis_beta.axhline(
            y=0,
            color="black",
            linewidth=1.0,
            linestyle="--",
            alpha=0.50
        )

        axis_beta.set_xlim(
            0,
            length
        )

        axis_beta.set_xlabel(
            "Distance from camera [m]",
            fontsize=11
        )

        axis_beta.set_ylabel(
            "Vertical viewing angle beta [deg]",
            fontsize=11
        )

        axis_beta.grid(
            True,
            alpha=0.25
        )

        # -------------------------------------------------
        # SECOND Y-AXIS FOR DELTA BETA
        # -------------------------------------------------

        axis_delta = axis_beta.twinx()

        if x_delta:

            line_delta, = axis_delta.plot(
                x_delta,
                delta_beta_values,
                color=COLOR_DELTA,
                linewidth=1.3,
                label="Delta beta (2230 - 2026)"
            )

            line_handles.append(
                line_delta
            )

        axis_delta.set_ylabel(
            "Delta beta: 2230 - 2026 [deg]",
            color=COLOR_DELTA,
            fontsize=11
        )

        axis_delta.tick_params(
            axis="y",
            labelcolor=COLOR_DELTA
        )

        axis_delta.axhline(
            y=0,
            color=COLOR_DELTA,
            linewidth=0.8,
            linestyle=":",
            alpha=0.50
        )

        axis_beta.set_title(
            f"Viewing-Angle Profile and Change - "
            f"PNumber {pnumber}",
            fontsize=14
        )

        if line_handles:

            labels = [
                line.get_label()
                for line in line_handles
            ]

            axis_beta.legend(
                line_handles,
                labels,
                loc="best"
            )

        fig.tight_layout()

        difference_plot_path = (
            plot_difference_dir
            / f"viewing_angle_change_{filename_id}.png"
        )

        fig.savefig(
            difference_plot_path,
            dpi=200,
            bbox_inches="tight"
        )

        plt.close(
            fig
        )

        # =================================================
        # CONSOLE SUMMARY
        # =================================================

        valid_beta_2026_count = sum(
            row["beta_2026_deg"] is not None
            for row in rows
        )

        valid_beta_2230_count = sum(
            row["beta_2230_deg"] is not None
            for row in rows
        )

        valid_delta_values = [
            row["delta_beta_2230_minus_2026_deg"]
            for row in rows
            if row[
                "delta_beta_2230_minus_2026_deg"
            ] is not None
        ]

        print(
            f"  Samples: {len(rows)}"
        )

        print(
            f"  Valid beta values in 2026: "
            f"{valid_beta_2026_count}"
        )

        print(
            f"  Valid beta values in 2230: "
            f"{valid_beta_2230_count}"
        )

        if valid_delta_values:

            print(
                "  Delta beta (2230 - 2026):"
            )

            print(
                f"    Minimum: "
                f"{min(valid_delta_values):.6f} deg"
            )

            print(
                f"    Maximum: "
                f"{max(valid_delta_values):.6f} deg"
            )

            print(
                f"    Mean:    "
                f"{sum(valid_delta_values) / len(valid_delta_values):.6f} deg"
            )

        print(
            f"  Excel: {excel_path.name}"
        )

        print(
            f"  Plot 1: {comparison_plot_path.name}"
        )

        print(
            f"  Plot 2: {difference_plot_path.name}"
        )

        print()

        processed_count += 1

    # =====================================================
    # FINISHED
    # =====================================================

    print("=" * 60)
    print("VIEWING-ANGLE PROCESSING COMPLETED")
    print("=" * 60)

    print(
        f"Processed lines: {processed_count}"
    )

    print(
        f"Skipped lines: {skipped_count}"
    )

    print(
        f"Output directory:\n{OUTPUT_DIR}"
    )

    print()

    print(
        "Definitions:"
    )

    print(
        "Beta = atan2(zTerrain - zEye, distance)"
    )

    print(
        "Delta Beta = Beta DEM 2230 - Beta DEM 2026"
    )

    print(
        "Positive Delta Beta: future viewing angle is higher."
    )

    print(
        "Negative Delta Beta: future viewing angle is lower."
    )

    print(
        f"Camera height above terrain: "
        f"{CAMERA_HEIGHT:.2f} m"
    )

    print(
        "Camera assumption: the observer remains 1.8 m "
        "above the terrain in each period."
    )

finally:

    qgs.exitQgis()